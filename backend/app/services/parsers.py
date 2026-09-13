import io
import os

def parse_document(file_bytes: bytes, filename: str) -> str:
    """
    Parses PDF or DOCX file content into plain text.
    Falls back gracefully if libraries or OCR are partially unavailable.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _parse_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        return _parse_docx(file_bytes)
    elif ext in [".txt", ".md"]:
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        # Fallback decoding
        return file_bytes.decode("utf-8", errors="ignore")

def _parse_pdf(file_bytes: bytes) -> str:
    text = ""
    # Try pdfplumber
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            extracted = []
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted.append(page_text)
            text = "\n".join(extracted)
            if text.strip():
                return text
    except Exception:
        pass

    # Try PyMuPDF (fitz)
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        extracted = []
        for page in doc:
            extracted.append(page.get_text())
        text = "\n".join(extracted)
        if text.strip():
            return text
    except Exception:
        pass

    # OCR Fallback via pytesseract if available
    try:
        import pytesseract
        from PIL import Image
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        ocr_text = []
        for page in doc:
            pix = page.get_pixmap()
            img = Image.open(io.BytesIO(pix.tobytes()))
            ocr_text.append(pytesseract.image_to_string(img))
        text = "\n".join(ocr_text)
        if text.strip():
            return text
    except Exception:
        pass

    # Ultimate fallback string extraction
    return file_bytes.decode("utf-8", errors="ignore")

def _parse_docx(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    full_text.append(row_text)
        return "\n".join(full_text)
    except Exception:
        return file_bytes.decode("utf-8", errors="ignore")
